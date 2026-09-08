"""공통 유틸: 뱅크샐러드 export 읽기, 프로필/룰 읽기, 가맹점 정규화.

의존: openpyxl, pyyaml (requirements.txt)
"""
from __future__ import annotations

import csv
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
    import yaml
except ImportError:  # pragma: no cover
    sys.exit("먼저 실행: python3 -m pip install -r requirements.txt")

import os
ROOT = Path(__file__).resolve().parent.parent
PRIVATE = Path(os.environ.get("CASHFLOW_PRIVATE", ROOT / "private"))   # 테스트용: 다른 폴더를 private 로
EXPORTS = PRIVATE / "exports"
SETUP = PRIVATE / "setup"
LEDGER = PRIVATE / "ledger"
OUT = PRIVATE / "out"
ENRICH = PRIVATE / "enrich"
TEMPLATES = ROOT / "templates"

TX_SHEET = "가계부 내역"
STATUS_SHEET = "뱅샐현황"
TX_COLUMNS = ["날짜", "시간", "타입", "대분류", "소분류", "내용", "금액", "화폐", "결제수단", "메모"]


# ---------------------------------------------------------------- 거래
@dataclass
class Tx:
    date: date
    time: str
    type: str          # 지출 | 수입 | 이체
    bs_cat: str        # 뱅샐 대분류
    bs_sub: str        # 뱅샐 소분류
    desc: str          # 내용(가맹점/상대)
    amount: int        # 부호 포함
    account: str       # 결제수단
    memo: str
    idx: int = 0

    @property
    def month(self) -> str:
        return self.date.strftime("%Y-%m")


def find_exports(paths: list[str] | None = None) -> list[Path]:
    if paths:
        return [Path(p) for p in paths]
    files = sorted((f for f in EXPORTS.rglob("*.xlsx") if not f.name.startswith("~$")), key=lambda p: p.stat().st_mtime)
    if not files:
        sys.exit(f"export 파일이 없습니다. 뱅크샐러드 '파일로 받기' xlsx를 {EXPORTS}/ 에 넣어주세요.")
    return files


def load_transactions(path: Path) -> list[Tx]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if TX_SHEET not in wb.sheetnames:
        sys.exit(f"{path.name}: '{TX_SHEET}' 시트가 없습니다. 뱅크샐러드 export 파일이 맞는지 확인하세요.")
    ws = wb[TX_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    missing = [c for c in TX_COLUMNS if c not in header]
    if missing:
        sys.exit(f"{path.name}: 기대 컬럼 누락 {missing}. 헤더={header}")
    col = {c: header.index(c) for c in TX_COLUMNS}
    out: list[Tx] = []
    for i, r in enumerate(rows):
        if r is None or r[col["날짜"]] is None:
            continue
        d = r[col["날짜"]]
        if isinstance(d, datetime):
            d = d.date()
        elif isinstance(d, str):
            d = datetime.strptime(d[:10], "%Y-%m-%d").date()
        t = r[col["시간"]]
        t = t.strftime("%H:%M") if hasattr(t, "strftime") else (str(t)[:5] if t else "")
        amt = r[col["금액"]]
        try:
            amt = int(round(float(amt)))
        except (TypeError, ValueError):
            continue
        out.append(Tx(
            date=d, time=t,
            type=str(r[col["타입"]] or "").strip(),
            bs_cat=str(r[col["대분류"]] or "미분류").strip(),
            bs_sub=str(r[col["소분류"]] or "미분류").strip(),
            desc=str(r[col["내용"]] or "").strip(),
            amount=amt,
            account=str(r[col["결제수단"]] or "").strip(),
            memo=str(r[col["메모"]] or "").strip(),
            idx=i,
        ))
    return out


def tx_key(t: Tx) -> str:
    """overrides.csv에서 거래 하나를 가리키는 키."""
    return f"{t.date.isoformat()}|{t.account}|{t.desc}|{t.amount}"


# ---------------------------------------------------------------- 현황 시트
@dataclass
class Status:
    person: dict = field(default_factory=dict)          # 이름/성별/연령/신용점수
    assets: list[dict] = field(default_factory=list)    # {group, name, value}
    liabilities: list[dict] = field(default_factory=list)
    insurance: list[dict] = field(default_factory=list)
    investments: list[dict] = field(default_factory=list)
    loans: list[dict] = field(default_factory=list)
    cashflow_bs: dict = field(default_factory=dict)     # 뱅샐 자체 월별 현금흐름 {항목: {월: 값}}
    exported_at: str = ""


def _cells(row) -> list:
    return [c for c in row]


def load_status(path: Path) -> Status:
    """'뱅샐현황' 시트를 섹션별로 파싱. 수식 셀은 무시(data_only)."""
    st = Status()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if STATUS_SHEET not in wb.sheetnames:
        return st
    m = re.search(r"~(\d{4}-\d{2}-\d{2})", path.name)
    st.exported_at = m.group(1) if m else ""
    rows = [list(r) for r in wb[STATUS_SHEET].iter_rows(values_only=True)]

    def section_bounds():
        marks = {}
        for i, r in enumerate(rows):
            for v in r:
                if isinstance(v, str):
                    mm = re.match(r"^(\d)\.(\S+)", v.strip())
                    if mm:
                        marks[int(mm.group(1))] = i
        order = sorted(marks.items())
        bounds = {}
        for k, (num, start) in enumerate(order):
            end = order[k + 1][1] if k + 1 < len(order) else len(rows)
            bounds[num] = (start, end)
        return bounds

    b = section_bounds()

    # 1. 고객정보
    if 1 in b:
        s, e = b[1]
        for i in range(s, e):
            vals = [v for v in rows[i] if v is not None]
            if vals and str(vals[0]).startswith("이름"):
                nxt = [v for v in rows[i + 1] if v is not None] if i + 1 < e else []
                keys = ["이름", "성별", "연령", "신용점수", "이메일"]
                st.person = {k: nxt[j] for j, k in enumerate(keys) if j < len(nxt)}
                break

    # 2. 현금흐름 (뱅샐 자체 집계) — 월 헤더 행 찾기
    if 2 in b:
        s, e = b[2]
        months, mcols = [], []
        for i in range(s, e):
            r = rows[i]
            if any(isinstance(v, str) and re.match(r"^\d{4}-\d{2}$", v) for v in r):
                for j, v in enumerate(r):
                    if isinstance(v, str) and re.match(r"^\d{4}-\d{2}$", v):
                        months.append(v); mcols.append(j)
                label_col = next(j for j, v in enumerate(r) if v == "항목")
                for k in range(i + 1, e):
                    rr = rows[k]
                    lab = rr[label_col]
                    if not isinstance(lab, str) or "총계" in lab:
                        continue
                    vals = {}
                    for mo, j in zip(months, mcols):
                        v = rr[j]
                        if isinstance(v, (int, float)):
                            vals[mo] = int(v)
                    if vals:
                        # 같은 라벨(미분류)이 수입/지출 양쪽에 있어 접미사로 구분
                        key = lab if lab not in st.cashflow_bs else lab + "(지출)"
                        st.cashflow_bs[key] = vals
                break

    # 3. 재무현황 (자산 | 부채)
    if 3 in b:
        s, e = b[3]
        hdr = None
        for i in range(s, e):
            r = rows[i]
            idx = [j for j, v in enumerate(r) if v == "상품명"]
            if len(idx) >= 1:
                hdr = i
                vals = [j for j, v in enumerate(r) if v == "금액"]   # 상품명이 병합셀이라 금액 위치를 따로 찾는다
                a_item, a_name = idx[0] - 1, idx[0]
                a_val = next(j for j in vals if j > a_name)
                if len(idx) >= 2:
                    l_item, l_name = idx[1] - 1, idx[1]
                    l_val = next(j for j in vals if j > l_name)
                else:
                    l_item = l_name = l_val = None
                break
        if hdr is not None:
            grp_a = grp_l = None
            for i in range(hdr + 1, e):
                r = rows[i]
                if r[a_item] and isinstance(r[a_item], str) and "총자산" in r[a_item]:
                    break
                if isinstance(r[a_item], str) and r[a_item].strip():
                    grp_a = r[a_item].strip()
                if r[a_name] is not None and isinstance(r[a_val], (int, float)):
                    st.assets.append({"group": grp_a, "name": str(r[a_name]).strip(), "value": float(r[a_val])})
                if l_item is not None:
                    if isinstance(r[l_item], str) and r[l_item].strip():
                        grp_l = r[l_item].strip()
                    if r[l_name] is not None and isinstance(r[l_val], (int, float)):
                        st.liabilities.append({"group": grp_l, "name": str(r[l_name]).strip(), "value": float(r[l_val])})

    def table(num, keys):
        if num not in b:
            return []
        s, e = b[num]
        out = []
        start = None
        for i in range(s, e):
            r = rows[i]
            if keys[0] in r and keys[1] in r:
                start = i; cols = [r.index(k) for k in keys]; break
        if start is None:
            return out
        for i in range(start + 1, e):
            r = rows[i]
            first = r[cols[0]]
            if first is None or (isinstance(first, str) and first.startswith("총")):
                if first is not None:
                    break
                continue
            out.append({k: r[j] for k, j in zip(keys, cols)})
        return out

    st.insurance = table(4, ["금융사", "보험명", "계약상태", "총납입금", "계약일자", "만기일자"])
    st.investments = table(5, ["투자상품종류", "금융사", "상품명", "투자원금", "평가금액", "수익률"])
    st.loans = table(6, ["대출종류", "금융사", "상품명", "대출원금", "대출잔액", "대출금리", "대출신규일", "대출만기일"])
    return st


# ---------------------------------------------------------------- 프로필 / 룰
def load_profile(path: Path | None = None) -> dict:
    import os
    p = path or Path(os.environ.get("CASHFLOW_PROFILE", PRIVATE / "profile.yaml"))
    if not p.exists():
        sys.exit(f"프로필이 없습니다: {p}\n  → scripts/inspect_export.py 로 초안을 만든 뒤 private/profile.yaml 로 저장하세요.")
    with open(p, encoding="utf-8") as f:
        prof = yaml.safe_load(f) or {}
    prof.setdefault("household", {}).setdefault("members", [])
    prof.setdefault("accounts", [])
    prof.setdefault("transfers", {})
    prof.setdefault("recurring", [])
    prof.setdefault("income", [])
    prof.setdefault("options", {})
    return prof


def load_rules(path: Path | None = None) -> list[dict]:
    """rules.csv: 키워드,카테고리,세부,고정비,귀속,메모 — 위에서부터 우선, 부분 포함 매칭."""
    p = path or PRIVATE / "rules.csv"
    if not p.exists():
        return []
    with open(p, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.DictReader(f) if r.get("키워드", "").strip() and not r["키워드"].startswith("#")]
    return rows


def load_overrides(path: Path | None = None) -> dict[str, dict]:
    """overrides.csv: 거래키,카테고리,세부,귀속,고정비,메모 — 거래 한 건에 대한 사용자 답."""
    p = path or LEDGER / "overrides.csv"
    if not p.exists():
        return {}
    with open(p, encoding="utf-8-sig", newline="") as f:
        return {r["거래키"]: r for r in csv.DictReader(f) if r.get("거래키")}


# ---------------------------------------------------------------- 정규화
_BRANCH = re.compile(r"(\s*\(.*?\)|\s+\S*점$|\s*\d+호?$|[\d\-\*]{4,})")


def normalize_merchant(desc: str) -> str:
    """'스타벅스 판교점' → '스타벅스', '이마트24 R판교' → '이마트24'. 인터뷰 묶음용."""
    d = desc.strip()
    d = re.sub(r"\(.*?\)", " ", d)
    d = re.sub(r"[\d\-\*]{4,}", " ", d)          # 카드번호·전화 조각
    d = re.sub(r"\s+", " ", d).strip()
    if not d:
        return desc
    tok = d.split(" ")
    head = tok[0]
    if len(head) <= 2 and len(tok) > 1:
        head = tok[0] + " " + tok[1]
    head = re.sub(r"(점|지점)$", "", head) if len(head) > 3 else head
    return head


def guess_kind(name: str) -> str:
    n = name
    if any(k in n for k in ("카드", "체크", "하이패스")):
        return "card"
    if any(k in n for k in ("페이", "머니", "포인트", "간편결제")):
        return "pay"
    if any(k in n for k in ("마이너스", "대출", "자금")):
        return "loan"
    if any(k in n for k in ("청약", "적금", "저금통", "박스", "모으기", "예금")) and "저축예금" not in n and "입출금" not in n:
        return "savings"
    if any(k in n for k in ("연금", "IRP", "퇴직")):
        return "pension"
    if any(k in n for k in ("증권", "위탁", "CMA", "ISA", "종합계좌", "옵션", "투자", "펀드")):
        return "investment"
    if any(k in n for k in ("보험",)):
        return "insurance"
    if "현금" == n:
        return "cash"
    return "checking"


def won(v) -> str:
    try:
        v = int(round(float(v)))
    except (TypeError, ValueError):
        return str(v)
    return f"{v:,}원"


# ---------------------------------------------------------------- 품목 보강 (쿠팡·네이버페이)
_DATE_RE = re.compile(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})")


def _parse_any_date(s: str):
    m = _DATE_RE.search(str(s or ""))
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _won_int(s) -> int | None:
    t = re.sub(r"[^\d\-]", "", str(s or ""))
    return int(t) if re.fullmatch(r"-?\d+", t) else None


def load_enrichment() -> dict[str, list[dict]]:
    """private/enrich/{coupang,naver}/<월>/*.csv → [{date, amount, text, src}]. 없으면 빈 목록."""
    out = {"coupang": [], "naver": [], "kurly": []}
    for f in sorted((ENRICH / "coupang").glob("*/쿠팡_주문_*.csv")):
        with open(f, encoding="utf-8-sig", newline="") as fh:
            for r in csv.DictReader(fh):
                d = _parse_any_date(r.get("주문일")); a = _won_int(r.get("총결제금액"))
                if d and a:
                    out["coupang"].append({"date": d, "amount": abs(a), "text": (r.get("상품요약") or r.get("구분") or "").strip(), "src": f.name, "used": False})
    for f in sorted((ENRICH / "kurly").glob("*/컬리_주문_*.csv")):
        with open(f, encoding="utf-8-sig", newline="") as fh:
            for r in csv.DictReader(fh):
                d = _parse_any_date(r.get("주문일")); a = _won_int(r.get("총결제금액"))
                if d and a:
                    out["kurly"].append({"date": d, "amount": abs(a), "text": (r.get("상품요약") or "").strip(), "src": f.name, "used": False})
    for f in sorted((ENRICH / "naver").glob("*/네이버페이_주문_*.csv")):
        with open(f, encoding="utf-8-sig", newline="") as fh:
            for r in csv.DictReader(fh):
                d = _parse_any_date(r.get("날짜")); a = _won_int(r.get("금액"))
                if d and a:
                    txt = " ".join(x for x in (r.get("가맹점", "").strip(), r.get("상품명", "").strip()) if x)
                    out["naver"].append({"date": d, "amount": abs(a), "text": txt, "status": r.get("상태", ""), "src": f.name, "used": False})
    return out


def match_enrichment(t: Tx, enrich: dict, days: int = 3) -> str:
    """거래 내용이 쿠팡/네이버페이면 같은 금액·근접 날짜의 주문을 찾아 품목 문자열을 돌려준다."""
    desc = (t.desc or "").lower()
    pool = None
    if "쿠팡" in desc or "coupang" in desc:
        pool = enrich["coupang"]
    elif "컬리" in desc or "kurly" in desc:
        pool = enrich["kurly"]
    elif "네이버페이" in desc or "naver" in desc or t.account.startswith("네이버페이"):
        pool = enrich["naver"]
    if not pool:
        return ""
    amt = abs(t.amount)
    best = None
    for o in pool:
        if o["used"] or o["amount"] != amt:
            continue
        dd = abs((o["date"] - t.date).days)
        if dd <= days and (best is None or dd < best[0]):
            best = (dd, o)
    if best:
        best[1]["used"] = True
        return best[1]["text"]
    return ""


# ---------------------------------------------------------------- AI 추정 (suggestions.csv)
def load_suggestions(path: Path | None = None, min_conf: float = 0.6) -> dict[str, dict]:
    """suggestions.csv: 묶음,카테고리,세부,신뢰도,근거,고정비,확정 — AI가 확인필요 묶음에 대해 추측한 것.
    신뢰도 < min_conf 는 버린다(결정적 신호를 확률로 덮지 않기 위해 룰·품목보다 항상 뒤)."""
    p = path or LEDGER / "suggestions.csv"
    if not p.exists():
        return {}
    out = {}
    with open(p, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            k = (r.get("묶음") or "").strip()
            if not k or k.startswith("#") or not r.get("카테고리"):
                continue
            try:
                conf = float(r.get("신뢰도") or 0)
            except ValueError:
                conf = 0.0
            if conf >= min_conf:
                r["_conf"] = conf
                out[k] = r
    return out
